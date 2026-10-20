import re
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
from utils import (cm_to_twips, twips_from_spacing, cn_size_to_pt, is_cjk, is_digit,
                   char_width_twips, is_heading_style, get_paragraph_text,
                   clear_paragraph_content, has_drawing_or_object,
                   set_line_spacing, set_space_before, set_space_after,
                   set_left_indent, set_right_indent, set_first_line_indent,
                   set_snap_to_grid, apply_style_to_run, apply_paragraph_spacing)


def clear_all_styles(doc):
    styles = doc.styles.element
    for style_el in list(styles.findall(qn('w:style'))):
        styles.remove(style_el)
    normal = OxmlElement('w:style')
    normal.set(qn('w:type'), 'paragraph')
    normal.set(qn('w:default'), '1')
    normal.set(qn('w:styleId'), 'Normal')
    name = OxmlElement('w:name')
    name.set(qn('w:val'), 'Normal')
    normal.append(name)
    styles.append(normal)


def clear_indents(doc):
    for p in doc.paragraphs:
        pPr = p._element.get_or_add_pPr()
        ind = pPr.find(qn('w:ind'))
        if ind is not None:
            for attr in ('left', 'right', 'firstLine', 'hanging', 'start', 'end'):
                if ind.get(qn('w:' + attr)) is not None:
                    ind.set(qn('w:' + attr), '0')


def clear_heading_indents(doc, clear_left, clear_right, clear_special):
    for p in doc.paragraphs:
        style_id = p.style.style_id if p.style else ''
        if not is_heading_style(style_id):
            continue
        pPr = p._element.get_or_add_pPr()
        ind = pPr.find(qn('w:ind'))
        if ind is None:
            continue
        if clear_left and ind.get(qn('w:left')) is not None:
            ind.set(qn('w:left'), '0')
        if clear_right and ind.get(qn('w:right')) is not None:
            ind.set(qn('w:right'), '0')
        if clear_special and ind.get(qn('w:firstLine')) is not None:
            ind.set(qn('w:firstLine'), '0')


def clear_snap_to_grid(doc):
    for p in doc.paragraphs:
        pPr = p._element.get_or_add_pPr()
        set_snap_to_grid(pPr, True)


def clear_extra_spaces(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            text = r.text
            if not text:
                continue
            cleaned = re.sub(r'[ \t]+', ' ', text)
            if cleaned != text:
                r.text = cleaned


def clear_chart_format(doc):
    pass


def flatten_sdt(doc):
    W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
    body = doc.element.body
    changed = True
    while changed:
        changed = False
        for sdt in list(body.iter('{%s}sdt' % W)):
            parent = sdt.getparent()
            if parent is None:
                continue
            idx = parent.index(sdt)
            sdt_content = sdt.find('{%s}sdtContent' % W)
            if sdt_content is not None:
                content_children = list(sdt_content)
                parent.remove(sdt)
                for i, child in enumerate(content_children):
                    parent.insert(idx + i, child)
            else:
                parent.remove(sdt)
            changed = True
            break


def replace_soft_returns(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            has_break = False
            for elem in list(r._element):
                if elem.tag.endswith('}br') and elem.get(qn('w:type')) == 'textWrapping':
                    has_break = True
                    r._element.remove(elem)
            if has_break and r.text:
                r._element.append(OxmlElement('w:br'))


def process_tab_mode(doc, mode):
    for p in doc.paragraphs:
        for r in p.runs:
            text = r.text
            if '\t' not in text:
                continue
            if mode == 1:
                r.text = text.replace('\t', ' ')
            elif mode == 2:
                r.text = text.replace('\t', '')
            elif mode == -1:
                r.text = text.replace('\t', '')
            elif mode == -2:
                r.text = text.replace('\t', ' ')


def insert_cjk_eu_spacing(doc, space_count=1):
    sp = ' ' * max(1, min(4, space_count))
    re1 = re.compile(r'([\u4e00-\u9fff\u3400-\u4dbf])([a-zA-Z0-9])')
    re2 = re.compile(r'([a-zA-Z0-9])([\u4e00-\u9fff\u3400-\u4dbf])')
    for p in doc.paragraphs:
        for r in p.runs:
            old = r.text
            if not old:
                continue
            new = re1.sub(r'\1' + sp + r'\2', old)
            new = re2.sub(r'\1' + sp + r'\2', new)
            if new != old:
                r.text = new


def markdown_to_plaintext(doc):
    replacements = {'**': '', '__': '', '~~': ''}
    for p in doc.paragraphs:
        for r in p.runs:
            text = r.text
            if not text:
                continue
            modified = text
            for old, new in replacements.items():
                modified = modified.replace(old, new)
            if modified != text:
                r.text = modified


def clear_superscript(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            rPr = r._element.find(qn('w:rPr'))
            if rPr is not None:
                vert_align = rPr.find(qn('w:vertAlign'))
                if vert_align is not None:
                    rPr.remove(vert_align)


def normalize_punctuation(doc):
    for p in doc.paragraphs:
        for r in p.runs:
            text = r.text
            if not text:
                continue
            out = normalize_punct(text)
            if out != text:
                r.text = out


def normalize_punct(text):
    chars = list(text)
    n = len(chars)
    result = []
    dq_state, sq_state = 0, 0
    for i, ch in enumerate(chars):
        prev = chars[i - 1] if i > 0 else ''
        nxt = chars[i + 1] if i + 1 < n else ''
        adj_cjk = is_cjk(prev) if prev else False
        if nxt:
            adj_cjk = adj_cjk or is_cjk(nxt)
        if ch == '.' and not is_digit(prev) and not is_digit(nxt) and adj_cjk:
            result.append('。')
            continue
        if ch == ',' and not is_digit(prev) and adj_cjk:
            result.append('，')
            continue
        if ch == ';' and adj_cjk:
            result.append('；')
            continue
        if ch == ':' and adj_cjk:
            result.append('：')
            continue
        if ch == '!' and adj_cjk:
            result.append('！')
            continue
        if ch == '?' and adj_cjk:
            result.append('？')
            continue
        if ch == '(' and adj_cjk:
            result.append('（')
            continue
        if ch == ')' and adj_cjk:
            result.append('）')
            continue
        if ch == '"' and adj_cjk:
            if dq_state == 0:
                result.append('“')
                dq_state = 1
            else:
                result.append('”')
                dq_state = 0
            continue
        if ch == "'" and adj_cjk:
            if sq_state == 0:
                result.append('‘')
                sq_state = 1
            else:
                result.append('’')
                sq_state = 0
            continue
        result.append(ch)
    return ''.join(result)


def clear_object_wrapping(doc):
    pass


def delete_empty_lines(doc):
    paras = doc.paragraphs
    to_remove = []
    for p in paras:
        text = get_paragraph_text(p)
        if text.strip() == '' and not has_drawing_or_object(p):
            to_remove.append(p)
    for p in to_remove:
        p._element.getparent().remove(p._element)


def adjust_caption_spacing(doc, cfg):
    re_fig = re.compile(r'^[图圖](\s*)([\d一二三四五六七八九十百千]+)')
    re_tbl = re.compile(r'^表(\s*)([\d一二三四五六七八九十百千]+)')
    for p in doc.paragraphs:
        text = get_paragraph_text(p)
        if re_fig.match(text) and cfg.get('preprocess_fig_caption', True):
            spaces = cfg.get('preprocess_fig_caption_spaces', 2) * 200
            pPr = p._element.get_or_add_pPr()
            set_space_before(pPr, spaces / 20, '磅')
            set_space_after(pPr, spaces / 20, '磅')
        elif re_tbl.match(text) and cfg.get('preprocess_tbl_caption', True):
            spaces = cfg.get('preprocess_tbl_caption_spaces', 2) * 200
            pPr = p._element.get_or_add_pPr()
            set_space_before(pPr, spaces / 20, '磅')
            set_space_after(pPr, spaces / 20, '磅')


def apply_preprocess_table_cells(doc, font, size_cn, line_spacing, align, border, min_height, auto_width):
    from format_table import apply_cell_borders
    from docx.enum.table import WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from utils import set_line_spacing, apply_style_to_run
    vmap = {'CENTER': WD_ALIGN_VERTICAL.CENTER, 'TOP': WD_ALIGN_VERTICAL.TOP, 'BOTTOM': WD_ALIGN_VERTICAL.BOTTOM}
    for tbl in doc.tables:
        if auto_width:
            tblPr = tbl._tbl.tblPr
            tblW = OxmlElement('w:tblW')
            tblW.set(qn('w:type'), 'auto')
            old = tblPr.find(qn('w:tblW'))
            if old is not None:
                tblPr.remove(old)
            tblPr.append(tblW)
            tblLayout = OxmlElement('w:tblLayout')
            tblLayout.set(qn('w:type'), 'autofit')
            old = tblPr.find(qn('w:tblLayout'))
            if old is not None:
                tblPr.remove(old)
            tblPr.append(tblLayout)
        for row in tbl.rows:
            if min_height and float(min_height) > 0:
                trPr = row._tr.get_or_add_trPr()
                trHeight = OxmlElement('w:trHeight')
                trHeight.set(qn('w:val'), str(int(min_height * 20)))
                trHeight.set(qn('w:hRule'), 'atLeast')
                trPr.append(trHeight)
            for cell in row.cells:
                if border:
                    apply_cell_borders(cell, border)
                if align and align in vmap:
                    cell.vertical_alignment = vmap[align]
                style = {'cn_font': font, 'en_font': font, 'size_cn': size_cn}
                if line_spacing and float(line_spacing) > 0:
                    style['line_spacing_mode'] = 'MULTIPLE'
                    style['line_spacing_value'] = float(line_spacing)
                for p in cell.paragraphs:
                    if line_spacing and float(line_spacing) > 0:
                        pPr = p._element.get_or_add_pPr()
                        set_line_spacing(pPr, 'MULTIPLE', float(line_spacing))
                    for r in p.runs:
                        apply_style_to_run(r, style)


def preprocess(doc, cfg):
    if cfg.get('preprocess_clear_all_styles'):
        clear_all_styles(doc)
    if cfg.get('preprocess_clear_indents'):
        clear_indents(doc)
    any_left = cfg.get('preprocess_clear_heading_indent') or cfg.get('preprocess_clear_heading_left_indent')
    any_right = cfg.get('preprocess_clear_heading_indent') or cfg.get('preprocess_clear_heading_right_indent')
    any_special = cfg.get('preprocess_clear_heading_indent') or cfg.get('preprocess_clear_heading_special_indent')
    if any_left or any_right or any_special:
        clear_heading_indents(doc, any_left, any_right, any_special)
    if cfg.get('preprocess_clear_snap_to_grid', True):
        clear_snap_to_grid(doc)
    if cfg.get('preprocess_clear_extra_spaces'):
        clear_extra_spaces(doc)
    if cfg.get('preprocess_clear_chart_format'):
        clear_chart_format(doc)
    if cfg.get('preprocess_flatten_sdt'):
        flatten_sdt(doc)
    if cfg.get('preprocess_soft_return'):
        replace_soft_returns(doc)
    tab_mode = cfg.get('preprocess_tab_mode', 0)
    if tab_mode and tab_mode != 0:
        process_tab_mode(doc, tab_mode)
    if cfg.get('preprocess_spaces'):
        body = cfg.get('body', {})
        if body.get('add_space'):
            insert_cjk_eu_spacing(doc, body.get('space_count', 1))
        else:
            insert_cjk_eu_spacing(doc, 1)
    if cfg.get('preprocess_markdown_to_plaintext'):
        markdown_to_plaintext(doc)
    if cfg.get('preprocess_superscript'):
        clear_superscript(doc)
    if cfg.get('preprocess_punctuation'):
        normalize_punctuation(doc)
    if cfg.get('preprocess_object_wrapping'):
        clear_object_wrapping(doc)
    if cfg.get('preprocess_delete_empty_lines'):
        delete_empty_lines(doc)
    adjust_caption_spacing(doc, cfg)
    if cfg.get('preprocess_table_cell_enable'):
        apply_preprocess_table_cells(doc,
            cfg.get('preprocess_table_cell_font', '宋体'),
            cfg.get('preprocess_table_cell_size', '四号'),
            cfg.get('preprocess_table_cell_line_spacing', 20),
            cfg.get('preprocess_table_cell_align', 'CENTER'),
            cfg.get('preprocess_table_cell_border', 'single'),
            cfg.get('preprocess_table_cell_min_height', 15),
            cfg.get('preprocess_table_auto_width', True))
