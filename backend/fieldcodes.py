from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
from utils import (cn_size_to_pt, get_paragraph_text, set_east_asia_font,
                   apply_style_to_run)


def _inject_toc_field_code(doc, toc_cfg, insert_after_title=False):
    levels = len(toc_cfg.get('level_styles', []))
    if levels < 1:
        levels = 4

    p = doc.add_paragraph()
    p.style = doc.styles['Normal']

    title_text = toc_cfg.get('title_text', '目录')
    title_cn = toc_cfg.get('title_cn_font', '')
    title_en = toc_cfg.get('title_en_font', '')
    title_size = toc_cfg.get('title_size_cn', '')

    title_run = p.add_run(title_text)
    rPr = title_run._element.get_or_add_rPr()
    if title_en or title_cn:
        font = title_en if title_en else title_cn
        title_run.font.name = font
        if title_cn and title_cn != title_en:
            set_east_asia_font(rPr, title_cn)
    if title_size:
        title_run.font.size = Pt(cn_size_to_pt(title_size))

    run_begin = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    run_begin._element.append(fld_begin)

    run_instr = p.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-%d" \\h \\z \\u' % levels
    run_instr._element.append(instr)

    run_sep = p.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    run_sep._element.append(fld_sep)

    placeholder_run = p.add_run('（打开文档后按 F9 刷新目录）')

    run_end = p.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run_end._element.append(fld_end)


def inject_toc(doc, toc_cfg):
    if not toc_cfg or not toc_cfg.get('enable'):
        return
    title_text = toc_cfg.get('title_text', '目录')
    found = False
    for p in doc.paragraphs:
        text = get_paragraph_text(p).strip()
        if text == title_text:
            found = True
            break
    _inject_toc_field_code(doc, toc_cfg, insert_after_title=found)
