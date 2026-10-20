from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from utils import (cn_size_to_pt, apply_style_to_run, set_color_hex,
                   set_east_asia_font, set_underline_type, ALIGN_MAP)


def _configure_run(run, cn_font, en_font, size_cn, bold, italic, underline_type):
    rPr = run._element.get_or_add_rPr()
    if en_font or cn_font:
        font = en_font if en_font else cn_font
        run.font.name = font
        if cn_font and cn_font != en_font:
            set_east_asia_font(rPr, cn_font)
    if size_cn:
        run.font.size = Pt(cn_size_to_pt(size_cn))
    if bold:
        run.font.bold = True
    if italic:
        run.font.italic = True
    if underline_type:
        set_underline_type(rPr, underline_type, '#000000')


def _add_field(paragraph, field_code):
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    run._element.append(fld_begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' %s ' % field_code
    instr_run._element.append(instr)

    sep_run = paragraph.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    sep_run._element.append(fld_sep)

    text_run = paragraph.add_run('1')

    end_run = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    end_run._element.append(fld_end)


def apply_header_footer(doc, hf):
    if not hf:
        return
    for section in doc.sections:
        if hf.get('enable_header'):
            header = section.header
            header.is_linked_to_previous = False
            for p in header.paragraphs:
                for r in p.runs:
                    r.text = ''
            if header.paragraphs:
                p = header.paragraphs[0]
                for r in p.runs:
                    r.text = ''
                align = hf.get('header_align', '')
                if align and align in ALIGN_MAP:
                    p.alignment = ALIGN_MAP[align]
                run = p.add_run()
                _configure_run(run, hf.get('header_cn_font', ''), hf.get('header_en_font', ''),
                               hf.get('header_size_cn', ''), hf.get('header_bold', False),
                               hf.get('header_italic', False), hf.get('header_underline_type', ''))
                run.text = hf.get('header_text', '')

        if hf.get('enable_footer'):
            footer = section.footer
            footer.is_linked_to_previous = False
            if hf.get('clear_footer', True):
                for p in footer.paragraphs:
                    for r in p.runs:
                        r.text = ''
            if footer.paragraphs:
                p = footer.paragraphs[0]
                for r in p.runs:
                    r.text = ''
                align = hf.get('footer_align', '')
                if align and align in ALIGN_MAP:
                    p.alignment = ALIGN_MAP[align]

                if not hf.get('page_number_from_body', False):
                    num_type = hf.get('footer_page_number_type', 'standard')
                    num_type_lower = num_type.lower() if num_type else ''
                    if num_type_lower in ('page_only', '仅页码'):
                        _add_field(p, 'PAGE')
                    elif num_type_lower in ('none', '无'):
                        pass
                    elif num_type_lower in ('standard', '标准'):
                        _add_field(p, 'PAGE')
                        p.add_run(' / ')
                        _add_field(p, 'NUMPAGES')
                    elif num_type_lower in ('dash', '短横线'):
                        _add_field(p, 'PAGE')
                        p.add_run(' - ')
                        _add_field(p, 'NUMPAGES')
                    elif num_type_lower in ('simple', '简单'):
                        _add_field(p, 'PAGE')
                    else:
                        _add_field(p, 'PAGE')
                        p.add_run(' / ')
                        _add_field(p, 'NUMPAGES')
